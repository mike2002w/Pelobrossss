#!/usr/bin/env python3
"""
Regenerate the workout tracker workbook from workouts.json.

The workbook is an OUTPUT now, not something to hand-edit: Sets and Cardio are
rewritten from the JSON, and the Progress / PR Summary formula grids are rebuilt
so new exercises and new sessions pick up formulas automatically (the old
workflow required dragging them across by hand).

Usage: python3 tools/json_to_xlsx.py workouts.json template.xlsx out.xlsx
"""
import json
import sys
from datetime import date, datetime

import openpyxl
from openpyxl.utils import get_column_letter

EQUIP_LABEL = {"Smith": "Smith", "DB": "Dumbbell", "Cable": "Cable",
               "Machine": "Machine", "BW": "Bodyweight"}
UNIT_LABEL = {"lb": "lbs", "kg": "kg", "bw": "BW"}
DAY = ["Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun"]

LAST = 2000  # formulas span Sets!$2:$2000, matching the original workbook


def day_of(iso):
    return DAY[date.fromisoformat(iso).weekday()]


def mmss(minutes):
    if minutes is None:
        return None
    total = int(round(float(minutes) * 60))
    return f"{total // 60}:{total % 60:02d}"


def clear(ws, first_row, last_row, last_col):
    for row in ws.iter_rows(min_row=first_row, max_row=last_row, max_col=last_col):
        for c in row:
            c.value = None


def write_sets(ws, workouts):
    clear(ws, 2, max(ws.max_row, LAST), 13)
    r = 2
    for w in workouts:
        for ex in w["exercises"]:
            for i, s in enumerate(ex["sets"], 1):
                ws.cell(r, 1, date.fromisoformat(w["date"]))
                ws.cell(r, 2, day_of(w["date"]))
                ws.cell(r, 3, w.get("typeLabel") or w["type"])
                ws.cell(r, 4, ex["sheetName"])
                ws.cell(r, 5, EQUIP_LABEL.get(ex["equipment"], ex["equipment"]))
                ws.cell(r, 6, UNIT_LABEL.get(ex["unit"], ex["unit"]))
                ws.cell(r, 7, i)
                ws.cell(r, 8, s["weight"])
                ws.cell(r, 9, s["reps"])
                ws.cell(r, 10, "Y" if s.get("warmup") else "N")
                ws.cell(r, 11, f'=IF($J{r}="Y","",IF($F{r}="BW",$I{r},$H{r}*$I{r}))')
                ws.cell(r, 12, f'=$D{r}&"|"&$H{r}')
                ws.cell(r, 13, ex.get("notes"))
                r += 1
    return r - 2


def write_cardio(ws, workouts, extra):
    bouts = [dict(b, date=w["date"]) for w in workouts if w.get("cardio")
             for b in [w["cardio"]]] + list(extra)
    bouts.sort(key=lambda b: (b["date"], b.get("machine") or ""))
    clear(ws, 2, max(ws.max_row, len(bouts) + 2), 11)
    for r, b in enumerate(bouts, start=2):
        ws.cell(r, 1, date.fromisoformat(b["date"]))
        ws.cell(r, 2, day_of(b["date"]))
        ws.cell(r, 3, b.get("machine"))
        ws.cell(r, 4, mmss(b.get("min")))
        ws.cell(r, 5, b.get("min"))
        ws.cell(r, 6, b.get("mi"))
        ws.cell(r, 7, "mi" if b.get("mi") is not None else None)
        ws.cell(r, 8, b.get("cal"))
        ws.cell(r, 9, b.get("hr"))
        ws.cell(r, 10, b.get("floors"))
        ws.cell(r, 11, b.get("notes"))
    return len(bouts)


def write_progress(ws, workouts, names):
    """One row per lifting session, one column per exercise, formulas filled in."""
    dates = [w["date"] for w in workouts if any(e["sets"] for e in w["exercises"])]
    clear(ws, 1, max(ws.max_row, len(dates) + 2), max(ws.max_column, len(names) + 2))
    ws.cell(1, 1, "Date")
    for i, n in enumerate(names, start=2):
        ws.cell(1, i, n)
    for r, d in enumerate(dates, start=2):
        ws.cell(r, 1, date.fromisoformat(d))
        for i in range(2, len(names) + 2):
            col = get_column_letter(i)
            crit = (f'Sets!$H$2:$H${LAST},Sets!$A$2:$A${LAST},$A{r},'
                    f'Sets!$D$2:$D${LAST},{col}$1,Sets!$J$2:$J${LAST},"N"')
            ws.cell(r, i, f'=IF(_xlfn.MAXIFS({crit})=0,"",_xlfn.MAXIFS({crit}))')
    return len(dates), len(names)


def write_pr(ws, meta, names, notes):
    clear(ws, 2, max(ws.max_row, len(names) + len(notes) + 4), 10)
    for r, n in enumerate(names, start=2):
        equip, unit = meta[n]
        ws.cell(r, 1, n)
        ws.cell(r, 2, EQUIP_LABEL.get(equip, equip))
        ws.cell(r, 3, UNIT_LABEL.get(unit, unit))
        ws.cell(r, 4, f'=_xlfn.MAXIFS(Sets!$H$2:$H${LAST},Sets!$D$2:$D${LAST},$A{r},Sets!$J$2:$J${LAST},"N")')
        ws.cell(r, 5, f'=IFERROR(INDEX(Sets!$I$2:$I${LAST},MATCH($A{r}&"|"&$D{r},Sets!$L$2:$L${LAST},0)),"")')
        ws.cell(r, 6, f'=IFERROR(INDEX(Sets!$A$2:$A${LAST},MATCH($A{r}&"|"&$D{r},Sets!$L$2:$L${LAST},0)),"")')
        ws.cell(r, 7, f'=_xlfn.MAXIFS(Sets!$A$2:$A${LAST},Sets!$D$2:$D${LAST},$A{r})')
        ws.cell(r, 8, f'=_xlfn.MAXIFS(Sets!$H$2:$H${LAST},Sets!$D$2:$D${LAST},$A{r},Sets!$A$2:$A${LAST},$G{r},Sets!$J$2:$J${LAST},"N")')
        ws.cell(r, 9, f'=COUNTIFS(Sets!$D$2:$D${LAST},$A{r},Sets!$J$2:$J${LAST},"N")')
        ws.cell(r, 10, f'=SUMIFS(Sets!$K$2:$K${LAST},Sets!$D$2:$D${LAST},$A{r},Sets!$J$2:$J${LAST},"N")')
    r = len(names) + 3
    for line in notes:
        ws.cell(r, 1, line)
        r += 1
    return len(names)


def main(json_path, template, out):
    data = json.load(open(json_path))
    workouts = data["workouts"]

    # stable ordering: workbook order is first-appearance, which keeps existing
    # Progress columns roughly where the user is used to seeing them
    names, meta = [], {}
    for w in workouts:
        for e in w["exercises"]:
            n = e["sheetName"]
            if n not in meta:
                meta[n] = (e["equipment"], e["unit"])
                names.append(n)

    wb = openpyxl.load_workbook(template)

    # preserve the hand-written notes block at the bottom of PR Summary
    pr = wb["PR Summary"]
    notes, seen = [], False
    for row in pr.iter_rows(min_row=2, max_col=1, values_only=True):
        if row[0] is None:
            continue
        if row[0] == "Notes":
            seen = True
        if seen:
            notes.append(row[0])

    n_sets = write_sets(wb["Sets"], workouts)
    n_cardio = write_cardio(wb["Cardio"], workouts, data.get("extraCardio", []))
    n_rows, n_cols = write_progress(wb["Progress"], workouts, names)
    n_pr = write_pr(pr, meta, names, notes)

    wb.save(out)
    print(f"  Sets:        {n_sets} rows")
    print(f"  Cardio:      {n_cardio} bouts")
    print(f"  Progress:    {n_rows} sessions x {n_cols} exercises")
    print(f"  PR Summary:  {n_pr} exercises (+{len(notes)} note lines kept)")
    print(f"  wrote {out}")


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2], sys.argv[3])
